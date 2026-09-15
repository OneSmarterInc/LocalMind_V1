"""Executable reference-adaptation tests; extraction uses real DOCX/PDF bytes."""
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from PIL import Image, ImageDraw
from docx import Document as Word
from docx.shared import Inches
from docx.oxml import OxmlElement
from .services.visuals import (AssetStore, extract_source_visuals, extraction_manifest,
                              filter_repeated_furniture, _save_pdf_visuals)
from .services.docx_visuals import normalized_picture
from .services.visual_context import choose_target, place_in_lesson


def picture_bytes():
    image = Image.new('RGB', (360, 220), 'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((15, 40, 135, 180), outline='black', width=3)
    draw.rectangle((220, 40, 345, 180), outline='black', width=3)
    draw.line((135, 105, 220, 105), fill='black', width=3)
    draw.text((28, 100), 'Sunlight', fill='black'); draw.text((245, 100), 'Leaf', fill='black')
    data = io.BytesIO(); image.save(data, format='PNG'); return data.getvalue()


def reference_docx(path):
    document = Word(); document.add_heading('Plant science', 1)
    for title, sentence, caption in (
        ('Sunlight', 'Chlorophyll absorbs sunlight in the leaves to supply energy for photosynthesis.', 'Figure 1: Sunlight and the leaf'),
        ('Water', 'Roots absorb water from the soil and transport it to leaves for photosynthesis.', 'Figure 2: Water and the leaf')):
        document.add_heading(title, 2); document.add_paragraph(sentence)
        document.add_picture(io.BytesIO(picture_bytes()), width=Inches(3))
        document.add_paragraph(caption, style='Caption')
    document.save(path)


class SourceVisualReferenceTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name)

    def test_docx_retains_context_caption_and_both_occurrences_of_same_image(self):
        source = self.folder/'source.docx'; reference_docx(source)
        rows = extract_source_visuals(source, self.folder/'processed')
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]['caption'], 'Figure 1: Sunlight and the leaf')
        self.assertEqual(rows[1]['heading_path'], ['Plant science', 'Water'])
        self.assertIn('Roots absorb water', rows[1]['context_text'])
        self.assertNotIn('Roots absorb water', rows[0]['context_text'])
        self.assertEqual(rows[0]['filename'], rows[1]['filename'])
        self.assertNotEqual(rows[0]['id'], rows[1]['id'])
        self.assertEqual(len(list((self.folder/'processed/visuals').glob('*.png'))), 1)
        self.assertEqual(rows[0]['width'], 360)

    def test_normalized_picture_honours_word_crop(self):
        raw, width, height = normalized_picture(picture_bytes(), (.5, 0, 0, 0))
        self.assertEqual((width, height), (180, 220))
        self.assertEqual(Image.open(io.BytesIO(raw)).format, 'PNG')

    def test_small_legitimate_compressed_figure_is_not_dropped_by_byte_size(self):
        self.assertLess(len(picture_bytes()), 6144)
        self.assertIsNotNone(normalized_picture(picture_bytes()))

    def test_thin_divider_is_excluded(self):
        out = io.BytesIO(); Image.new('RGB', (1000, 40), 'white').save(out, format='PNG')
        self.assertIsNone(normalized_picture(out.getvalue()))

    def test_invalid_or_active_image_is_rejected(self):
        with self.assertRaises(OSError): normalized_picture(b'<svg onload="alert(1)"/>')

    def test_external_relationship_is_not_followed_and_other_pictures_survive(self):
        source = self.folder/'source.docx'; reference_docx(source)
        with ZipFile(source) as archive: files = {n: archive.read(n) for n in archive.namelist()}
        relpath = 'word/_rels/document.xml.rels'
        import xml.etree.ElementTree as ET
        root = ET.fromstring(files[relpath])
        for rel in root:
            if rel.get('Type','').endswith('/image'):
                rel.set('TargetMode','External'); rel.set('Target','https://not-contacted.invalid/picture.png')
        files[relpath] = ET.tostring(root)
        with ZipFile(source, 'w') as archive:
            for n, raw in files.items(): archive.writestr(n, raw)
        self.assertEqual(extract_source_visuals(source, self.folder/'processed'), [])
        self.assertIn('external', extraction_manifest(self.folder/'processed')['warnings'][0])

    def test_cached_extraction_does_not_decode_again(self):
        source = self.folder/'source.docx'; reference_docx(source)
        rows = extract_source_visuals(source, self.folder/'processed')
        with patch('documents.services.docx_visuals.read_docx_visuals', side_effect=AssertionError('decoded twice')):
            self.assertEqual(extract_source_visuals(source, self.folder/'processed'), rows)

    def test_missing_cached_asset_is_rebuilt(self):
        source = self.folder/'source.docx'; reference_docx(source)
        rows = extract_source_visuals(source, self.folder/'processed')
        (self.folder/'processed/visuals'/rows[0]['filename']).unlink()
        self.assertEqual(len(extract_source_visuals(source, self.folder/'processed')), 2)
        self.assertTrue((self.folder/'processed/visuals'/rows[0]['filename']).exists())

    def test_repeated_body_figure_is_retained_on_every_page(self):
        rows = [{'digest':'same','page':n,'margin_art':False} for n in range(1,6)]
        self.assertEqual(filter_repeated_furniture(rows), rows)

    def test_four_repeats_on_one_page_are_not_a_running_header(self):
        rows = [{'digest':'same','page':1,'margin_art':True} for _ in range(4)]
        self.assertEqual(filter_repeated_furniture(rows), rows)

    def test_repeated_margin_logo_does_not_remove_a_body_occurrence(self):
        rows = [{'digest':'same','page':n,'margin_art':True} for n in range(1,5)]
        rows.append({'digest':'same','page':5,'margin_art':False})
        self.assertEqual(filter_repeated_furniture(rows), rows[-1:])

    def test_context_disambiguates_modules_on_the_same_page(self):
        targets = [{'id':'water','title':'Water','source':'Roots absorb water from the soil and transport it to leaves.','start_page':1},
                   {'id':'sun','title':'Sunlight','source':'Chlorophyll absorbs sunlight in the leaves to supply energy.','start_page':1}]
        target, _ = choose_target({'page':1,'context_text':targets[1]['source']}, targets)
        self.assertEqual(target['id'], 'sun')

    def test_ambiguous_context_does_not_fall_back_to_first_module(self):
        targets = [{'id':i,'source':'The same repeated passage about carbon and water.','title':'Same'} for i in (1,2)]
        self.assertIsNone(choose_target({'context_text':targets[0]['source']}, targets)[0])

    def test_shared_chapter_and_generic_words_do_not_relocate_orphaned_picture(self):
        visual={'context_text':'Chlorophyll absorbs sunlight in the leaves to supply energy for photosynthesis.',
                'caption':'Figure 1: Sunlight and the leaf', 'caption_origin':'source',
                'heading_path':['Plant science','Sunlight']}
        targets=[{'id':'new','title':'New topic','chapter_title':'Plant science',
                  'source':'This separate topic explains foreign currency markets and international banking.'},
                 {'id':'water','title':'Water','chapter_title':'Plant science',
                  'source':'Roots absorb water from the soil and transport it to leaves for photosynthesis. Figure 2: Water and the leaf'}]
        self.assertIsNone(choose_target(visual, targets)[0])

    def test_known_page_never_borrows_another_pages_module(self):
        self.assertIsNone(choose_target({'page':9,'context_text':'Water moves through roots.'},
                         [{'id':1,'source':'Water moves through roots.','start_page':1,'end_page':2}])[0])

    def test_lesson_placement_uses_source_context_and_keeps_unknown_in_gallery(self):
        lesson = {'sections':[{'heading':'Water', 'explanation':'Roots absorb water from the soil and move it into the plant.', 'source_reference':'Roots absorb water from the soil and move it into the plant.'},
                              {'heading':'Energy', 'explanation':'Chlorophyll uses sunlight to supply energy for photosynthesis.', 'source_reference':'Chlorophyll uses sunlight to supply energy for photosynthesis.'}]}
        rows = [{'id':'known','context_text':lesson['sections'][1]['source_reference']}, {'id':'unknown','context_text':''}]
        result = place_in_lesson(lesson, rows)
        self.assertEqual(result['sections'][1]['visual_ids'], ['known'])
        self.assertEqual(result['sections'][0]['visual_ids'], [])
        self.assertEqual(len(result['source_visuals']), 2)
        self.assertNotIn('visual_ids', lesson['sections'][0])

    def test_rotated_pdf_figure_is_cropped_in_correct_coordinates(self):
        import pymupdf
        source = self.folder/'rotated.pdf'
        with pymupdf.open() as doc:
            page=doc.new_page(width=600,height=800)
            page.insert_image(pymupdf.Rect(90,230,450,450),stream=picture_bytes())
            page.insert_text((60,80),'This paragraph must not be included in the picture.')
            page.set_rotation(90); doc.save(source)
        rows = _save_pdf_visuals(source, self.folder/'png')
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]['bbox'], [90,230,450,450])
        self.assertEqual((rows[0]['width'],rows[0]['height']), (720,440))

    def test_page_sized_docx_picture_is_excluded(self):
        source = self.folder/'page.docx'; document = Word(); document.add_heading('Scan',1)
        document.add_picture(io.BytesIO(picture_bytes()), width=Inches(8), height=Inches(10.5));document.save(source)
        self.assertEqual(extract_source_visuals(source,self.folder/'processed'), [])
        self.assertIn('page-sized',extraction_manifest(self.folder/'processed')['warnings'][0])

if __name__ == '__main__': unittest.main()
