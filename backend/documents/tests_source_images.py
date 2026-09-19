import base64
import io
import json
import tempfile
from pathlib import Path
from unittest.mock import patch

from django.core.management import call_command
from django.test import TestCase, override_settings
from PIL import Image

from core.testing import make_faculty, make_student, make_subject, assign, enroll, make_published_document, client_for
from documents.services.visual_context import choose_target, place_in_lesson
from documents.services.visual_delivery import module_visuals
from documents.services.visuals import extract_source_visuals
from tutor.models import ModuleLesson
from tutor.lessons import source_hash


class SourceImageTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.settings = override_settings(MEDIA_ROOT=self.temp.name)
        self.settings.enable()
        self.addCleanup(self.settings.disable)
        self.subject = make_subject()
        self.doc = make_published_document(self.subject)
        self.module = self.doc.chapters.first().modules.first()
        self.module.start_page = self.module.end_page = 1
        self.module.save()
        self.base = Path(self.temp.name) / 'processed' / str(self.doc.pk)
        (self.base / 'visuals').mkdir(parents=True)
        self.doc.processed_markdown_path = str(self.base / 'source.md')
        self.doc.save()
        image = Image.new('RGB', (180, 100), 'blue')
        out = io.BytesIO(); image.save(out, format='PNG'); self.raw = out.getvalue()
        (self.base / 'visuals' / 'figure.png').write_bytes(self.raw)
        self.visual = dict(id='p1-figure', filename='figure.png', kind='figure', page=1,
                           caption='Figure 1: Processes', caption_origin='source', width=180, height=100,
                           context_text=self.module.source_text, heading_path=[])
        self.manifest([self.visual])
        self.lesson = dict(title='Processes', learning_objectives=[], key_terms=[], summary='Running programs',
                           sections=[dict(heading='Processes', explanation=self.module.source_text, source_reference=self.module.source_text)])
        ModuleLesson.objects.create(module=self.module,status='ready',source_hash=source_hash(self.module.source_text),lesson=self.lesson)

    def manifest(self, rows):
        (self.base / 'visuals' / 'manifest.json').write_text(json.dumps({'visuals': rows}))

    def test_authorized_faculty_student_and_offline_get_same_original_pixels(self):
        faculty = make_faculty(); assign(faculty, self.subject)
        snapshot = client_for(faculty).get(f'/api/faculty/modules/{self.module.pk}/local-authoring/').data
        student = make_student(); enroll(student,self.subject)
        client = client_for(student)
        path = f'/student/modules/{self.module.pk}/teach/'
        with patch('ai.gateway.gateway', side_effect=AssertionError('No inference for imagery')):
            lesson = client.get('/api'+path).data['lesson']
            bundle = client.get('/api/student/offline/').data['entries'][path]['lesson']
        read = client.get(f'/api/student/modules/{self.module.pk}/').data
        offline_read = client.get('/api/student/offline/').data['entries'][f'/student/modules/{self.module.pk}/']
        for delivered in [snapshot['source_visuals'], lesson['source_visuals'], bundle['source_visuals'], read['source_visuals'], offline_read['source_visuals']]:
            self.assertEqual(base64.b64decode(delivered[0]['data_url'].split(',')[1]),self.raw)
        self.assertEqual(lesson['sections'][0]['visual_ids'],['p1-figure'])
        outsider = make_student(email='outsider-image@example.com')
        self.assertIn(client_for(outsider).get('/api'+path).status_code,[403,404])
        self.assertEqual(ModuleLesson.objects.get(module=self.module).lesson,self.lesson)

    def test_ambiguous_module_and_path_escape_never_delivered(self):
        from learning.models import Module
        Module.objects.create(chapter=self.module.chapter,title=self.module.title,source_text=self.module.source_text,order=99,start_page=1,end_page=1)
        self.assertEqual(module_visuals(self.module),[])
        self.visual['filename']='../../outside.png'; self.manifest([self.visual])
        self.assertEqual(module_visuals(self.module),[])

    def test_unrelated_or_tied_figures_do_not_attach_to_first_section(self):
        lesson={**self.lesson,'sections':self.lesson['sections']*2}
        result=place_in_lesson(lesson,[self.visual])
        self.assertEqual([s['visual_ids'] for s in result['sections']],[[],[]])
        self.assertEqual(len(result['source_visuals']),1)
        self.assertNotIn('visual_ids',lesson['sections'][0])
        self.assertIsNone(choose_target({'page':8},[{'id':1,'start_page':1,'end_page':2}])[0])

    def test_docx_extraction_keeps_image_pixels_and_module_outline(self):
        from docx import Document as WordDocument
        doc=WordDocument();doc.add_heading(self.module.title,level=1)
        doc.add_paragraph(self.module.source_text)
        doc.add_picture(io.BytesIO(self.raw));doc.add_paragraph('Figure 1: Processes')
        path=Path(self.temp.name)/'book.docx';doc.save(path)
        self.doc.file.name='book.docx';self.doc.save()
        before=list(self.doc.chapters.values_list('id','title','source_text'))
        modules=list(self.doc.chapters.first().modules.values_list('id','title','source_text'))
        call_command('refresh_source_visuals',document=str(self.doc.pk),stdout=io.StringIO())
        rows=extract_source_visuals(path,self.base)
        self.assertEqual(len(rows),1)
        restored=Image.open(self.base/'visuals'/rows[0]['filename'])
        self.assertEqual(restored.tobytes(),Image.open(io.BytesIO(self.raw)).tobytes())
        self.assertEqual(before,list(self.doc.chapters.values_list('id','title','source_text')))
        self.assertEqual(modules,list(self.doc.chapters.first().modules.values_list('id','title','source_text')))
        self.assertEqual(ModuleLesson.objects.get(module=self.module).lesson,self.lesson)

    def test_pdf_raster_and_vector_figures_are_separate_from_outline(self):
        import pymupdf
        doc=pymupdf.open();page=doc.new_page(width=600,height=800)
        page.insert_text((50,50),'Processes and program execution')
        page.insert_image(pymupdf.Rect(80,100,380,280),stream=self.raw)
        page.insert_text((80,305),'Figure 1: Process states and transitions')
        page=doc.new_page(width=600,height=800)
        page.insert_text((50,50),'Vector diagram of process transitions')
        for x in (80,210,340):
            page.draw_rect(pymupdf.Rect(x,120,x+100,210),color=(0,0,0),fill=(.8,.9,.8))
            page.insert_text((x+12,165),'Process')
        page.draw_line(pymupdf.Point(180,165),pymupdf.Point(210,165))
        page.draw_line(pymupdf.Point(310,165),pymupdf.Point(340,165))
        page.insert_text((80,240),'Figure 2: Process transitions')
        path=Path(self.temp.name)/'book.pdf';doc.save(path);doc.close()
        rows=extract_source_visuals(path,self.base)
        self.assertTrue(rows)
        self.assertEqual({v['page'] for v in rows},{1,2})
        self.assertTrue(all(v['width']<1200 and v['height']<1600 for v in rows))
        self.assertEqual(rows,extract_source_visuals(path,self.base))


class StaffPictureTabTests(TestCase):
    """The read-only Pictures tab and its endpoints: every extracted picture,
    reachable from the book itself rather than only from a lesson."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        s = override_settings(MEDIA_ROOT=self.temp.name); s.enable(); self.addCleanup(s.disable)
        self.subject = make_subject()
        self.doc = make_published_document(self.subject)
        self.modules = list(self.doc.chapters.first().modules.all())
        for page, module in enumerate(self.modules, start=1):
            module.start_page = module.end_page = page
            module.save()
        self.base = Path(self.temp.name) / 'processed' / str(self.doc.pk)
        (self.base / 'visuals').mkdir(parents=True)
        self.doc.processed_markdown_path = str(self.base / 'source.md')
        self.doc.save()
        img = Image.new('RGB', (180, 100), 'blue')
        out = io.BytesIO(); img.save(out, format='PNG')
        (self.base / 'visuals' / 'a.png').write_bytes(out.getvalue())
        (self.base / 'visuals' / 'b.png').write_bytes(out.getvalue())
        visuals = [
            dict(id='p1-a', filename='a.png', kind='figure', page=1, caption='Figure 1',
                 caption_origin='source', width=180, height=100, context_text=self.modules[0].source_text, heading_path=[]),
            dict(id='p9-b', filename='b.png', kind='figure', page=9, caption='Orphan',
                 caption_origin='source', width=180, height=100, context_text='unrelated text', heading_path=[]),
        ]
        (self.base / 'visuals' / 'manifest.json').write_text(json.dumps({'visuals': visuals, 'warnings': []}))

    def test_index_counts_by_module_without_image_bytes(self):
        faculty = make_faculty(); assign(faculty, self.subject)
        res = client_for(faculty).get(f'/api/faculty/documents/{self.doc.pk}/pictures/')
        self.assertEqual(res.status_code, 200, res.data)
        self.assertEqual(res.data['total'], 2)
        self.assertEqual(res.data['assigned'], 1)
        self.assertEqual(res.data['needs_review'], 1)
        self.assertNotIn('data:image/png', str(res.data['chapters']))
        placed = [m for c in res.data['chapters'] for m in c['modules'] if m['count']]
        self.assertEqual(placed[0]['id'], str(self.modules[0].pk))

    def test_review_queue_carries_the_orphan_with_its_image_and_reason(self):
        faculty = make_faculty(); assign(faculty, self.subject)
        res = client_for(faculty).get(f'/api/faculty/documents/{self.doc.pk}/pictures/')
        self.assertEqual(len(res.data['review']), 1)
        row = res.data['review'][0]
        self.assertTrue(row['data_url'].startswith('data:image/png;base64,'))
        self.assertEqual(row['reason'], 'no_matching_source_page')

    def test_module_visuals_endpoint_returns_the_placed_picture(self):
        faculty = make_faculty(); assign(faculty, self.subject)
        res = client_for(faculty).get(f'/api/faculty/modules/{self.modules[0].pk}/visuals/')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(len(res.data['visuals']), 1)
        self.assertTrue(res.data['visuals'][0]['data_url'].startswith('data:image/png;base64,'))

    def test_a_student_cannot_reach_the_picture_index(self):
        student = make_student()
        self.assertIn(client_for(student).get(f'/api/faculty/documents/{self.doc.pk}/pictures/').status_code, (403, 404))
