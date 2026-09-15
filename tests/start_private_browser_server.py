"""Disposable REAL Django server for browser tests. Never uses the user's DB.

Only inference is stubbed in the normal browser suite. Set LM_E2E_REAL_MODEL=1
for the separate smoke that keeps the real bundled browser inference runtime.
"""
import json, os, shutil, subprocess, sys, tempfile
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1]
if os.environ.get('RUN_LOCALMIND_BROWSER_TESTS')!='1':
    raise SystemExit('Test server requires RUN_LOCALMIND_BROWSER_TESTS=1.')
with tempfile.TemporaryDirectory(prefix='localmind-browser-') as folder:
    work=Path(folder);web=work/'web'
    shutil.copytree(ROOT/'frontend/dist',web)
    if os.environ.get('LM_E2E_REAL_MODEL')!='1':
        shutil.copyfile(ROOT/'tests/fixtures/private-model-stub.js',web/'private-assets/runtime-loader.js')
    os.environ.update(DJANGO_SETTINGS_MODULE='config.settings',DJANGO_DEBUG='true',
        DJANGO_SECRET_KEY='localmind-disposable-browser-test-only-secret',
        DATABASE_URL=f'sqlite:///{work / "db.sqlite3"}',WEB_DIST=str(web),
        AI_ENABLED='false',AI_MONITOR_MODE='off',DURABLE_JOBS='false',
        AUTO_GENERATE_LESSONS='false',AUTO_GENERATE_QUIZZES='false',API_DOCS_ENABLED='false')
    sys.path.insert(0,str(ROOT/'backend'))
    import django
    django.setup()
    from django.conf import settings
    settings.MEDIA_ROOT=work/'media'
    settings.PRIVATE_LIBRARY_ROOT=work/'private-books'
    # Every callable sees these settings in this same process; no user .env is changed.
    from django.core.management import call_command
    call_command('migrate',interactive=False,verbosity=0)
    from core.testing import make_admin,make_faculty,make_student,make_subject,assign,enroll,make_published_document
    from django.core.files.base import ContentFile
    import io,hashlib
    from docx import Document as Word
    from reportlab.pdfgen.canvas import Canvas
    password='Disposable-Browser-Test-2026!'
    admin=make_admin(email='browser-admin@example.edu',password=password)
    faculty=make_faculty(email='browser-faculty@example.edu',password=password)
    student=make_student(email='browser-student@example.edu',password=password)
    outsider=make_student(email='browser-other@example.edu',password=password)
    subject=make_subject(code='WEBTEST',name='Browser test biology')
    assign(faculty,subject);enroll(student,subject)
    source=('Photosynthesis happens in the chloroplasts of green leaves. Chlorophyll absorbs sunlight. '
            'Plants take in carbon dioxide through stomata and release oxygen. Water enters the plant through roots. '
            'Sunlight supplies energy to produce glucose. Glucose can be stored as starch. ')*5
    doc=make_published_document(subject,title='Faculty Biology',modules=(('Leaf science',source),('Open practice',source)))
    out=io.BytesIO();word=Word();word.add_heading('Leaf science',1);word.add_paragraph(source)
    word.add_heading('Open practice',1);word.add_paragraph(source);word.save(out);raw=out.getvalue()
    doc.original_name='faculty-biology.docx';doc.file_type='docx';doc.file_size=len(raw);doc.content_hash=hashlib.sha256(raw).hexdigest();doc.uploaded_by=faculty
    doc.file.save(doc.original_name,ContentFile(raw),save=True)
    from learning.models import Module
    module=Module.objects.filter(chapter__document=doc).order_by('order').first()
    from assessments.models import Assessment
    from core.testing import MCQ
    question={**MCQ,'id':'course-q1'}
    immediate=Assessment.objects.create(subject=subject,module=module,kind='module',title='Offline immediate quiz',questions=[question],status='published',results_release='immediate')
    held=Assessment.objects.create(subject=subject,module=module,kind='module',title='Offline held quiz',questions=[question],status='published',results_release='held')
    from tutor.models import ModuleLesson
    from tutor.lessons import source_hash
    ModuleLesson.objects.update_or_create(module=module,defaults={'status':'ready','source_hash':source_hash(module.source_text),'lesson':{'title':'Stored course lesson','learning_objectives':['Understand plants'],'sections':[{'heading':'Photosynthesis','explanation':'This lesson was prepared by the institution before download.','source_reference':'Leaf science'}],'key_terms':[],'summary':'Plants use sunlight.'}})
    results=ROOT/'frontend/test-results';results.mkdir(exist_ok=True)
    pdf=io.BytesIO();canvas=Canvas(pdf);canvas.drawString(40,760,'Photosynthesis happens in the chloroplasts of green leaves.');canvas.save()
    (results/'private-fixture.pdf').write_bytes(pdf.getvalue())
    # Real image-only scan: source text, numeric table and diagram are rasterised first.
    from reportlab.lib.utils import ImageReader
    import pypdfium2 as pdfium
    source_pdf=io.BytesIO();c=Canvas(source_pdf,pagesize=(600,780))
    c.setFont('Helvetica-Bold',20);c.drawString(40,730,'Plant science')
    c.setFont('Helvetica',13)
    c.drawString(40,690,'Photosynthesis happens in the chloroplasts of green leaves.')
    c.drawString(40,670,'Chlorophyll absorbs sunlight. Roots absorb water from soil.')
    c.drawString(40,650,'Plants use carbon dioxide and water to produce glucose.')
    c.drawString(40,630,'Oxygen is released through stomata in the leaves.')
    c.drawString(40,590,'Measurements from the source table')
    for y in (560,520,480):c.line(40,y,550,y)
    for x in (40,300,550):c.line(x,480,x,560)
    c.drawString(50,535,'Plant');c.drawString(310,535,'Height in cm')
    c.drawString(50,495,'Bean');c.drawString(310,495,'12.5')
    c.rect(60,330,140,70);c.rect(350,330,140,70)
    c.drawString(85,360,'Sunlight');c.drawString(380,360,'Leaf')
    c.line(200,365,350,365);c.line(350,365,335,375);c.line(350,365,335,355)
    c.save(); original=pdfium.PdfDocument(source_pdf.getvalue());page=original[0]
    bitmap=page.render(scale=2);raster=bitmap.to_pil();raster.save(results/'scan-source.png')
    scan=io.BytesIO();c=Canvas(scan,pagesize=(600,780));c.drawImage(ImageReader(raster),0,0,600,780);c.save()
    (results/'scanned-biology.pdf').write_bytes(scan.getvalue())
    illustrated=io.BytesIO();c=Canvas(illustrated,pagesize=(600,780));c.drawImage(ImageReader(raster),350,30,200,260)
    for line in range(12):c.drawString(30,750-line*20,'Readable textbook content about sunlight, leaves and photosynthesis.')
    c.save();(results/'illustrated-text.pdf').write_bytes(illustrated.getvalue())
    # A small PDF whose expanded PNG pages exceed the former 48 MiB limit.
    import random
    from PIL import Image
    noise=Image.frombytes('RGB',(700,700),random.Random(42).randbytes(700*700*3))
    large=io.BytesIO();c=Canvas(large,pagesize=(600,780))
    for n in range(44):
        c.drawImage(ImageReader(noise),30,30,540,540)
        for line in range(8):c.drawString(30,750-line*18,'Readable textbook content about sunlight, leaves and photosynthesis.')
        c.showPage()
    c.save();(results/'large-illustrated.pdf').write_bytes(large.getvalue())
    bitmap.close();page.close();original.close()
    (results/'fixture.json').write_text(json.dumps({'quizImmediate':str(immediate.pk),'quizHeld':str(held.pk),'module':str(module.id),'document':str(doc.id),'subject':str(subject.id),'source':source,'password':password}))
    # runserver stays in this process so temporary storage settings are retained.
    call_command('runserver','127.0.0.1:8765',use_reloader=False,verbosity=0)
